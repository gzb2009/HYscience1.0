import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[4]
SCRIPT = ROOT / "backend/cli/skills/biology/scanpy/scripts/annotation_report.py"
FIXTURES = ROOT / "backend/cli/test/fixtures/annotation"
SPEC = importlib.util.spec_from_file_location("annotation_report", SCRIPT)
REPORT = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(REPORT)


class AnnotationReportTest(unittest.TestCase):
    def fixture(self):
        return json.loads((FIXTURES / "pdac_tnk_annotations.json").read_text())

    def test_normalizes_markers_confidence_and_complete_cluster_set(self):
        expected = REPORT.source_cluster_ids(str(FIXTURES / "pdac_tnk_markers.csv"))
        data, audit = REPORT.normalize_data(self.fixture(), expected)
        self.assertEqual(audit["expected_clusters"], 5)
        self.assertEqual(audit["written_clusters"], 5)
        self.assertEqual(data["clusters"][1]["confidence"], "⚠️")
        self.assertEqual(data["clusters"][0]["markers"], "Foxp3(66%), Il2ra(67%)")
        self.assertEqual(data["clusters"][1]["markers"], "Nkg7(98%), Prf1(83%)")

    def test_rejects_missing_duplicate_and_incomplete_clusters(self):
        data = self.fixture()
        data["clusters"] = data["clusters"][:-1]
        with self.assertRaisesRegex(REPORT.ValidationError, "missing cluster ids"):
            REPORT.normalize_data(data, ["0", "1", "2", "3", "4"])

        data = self.fixture()
        data["clusters"][1]["id"] = "C0"
        with self.assertRaisesRegex(REPORT.ValidationError, "duplicate cluster ids"):
            REPORT.normalize_data(data)

        data = self.fixture()
        data["clusters"][0]["notes"] = ""
        with self.assertRaisesRegex(REPORT.ValidationError, "missing: notes"):
            REPORT.normalize_data(data)

    def test_writes_four_consistent_readable_sheets(self):
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            self.skipTest("openpyxl is required for styled workbook verification")

        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            data = self.fixture()
            data["clusters"][0]["notes"] = "长文本证据。" * 80
            source = folder / "annotations.json"
            source.write_text(json.dumps(data, ensure_ascii=False))
            output = folder / "pdac_tnk_annotation.xlsx"
            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--input",
                    str(source),
                    "--markers",
                    str(FIXTURES / "pdac_tnk_markers.csv"),
                    "--output",
                    str(output),
                    "--locale",
                    "zh",
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("stats.total=99", result.stderr)

            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["TNK亚群汇总", "下游分析保留列表", "污染报告", "关键发现与建议"])
            summary = workbook["TNK亚群汇总"]
            self.assertEqual(summary.max_row, 6)
            self.assertEqual(summary["E3"].value, "⚠️")
            self.assertGreater(summary.row_dimensions[2].height, 36)

            contamination = workbook["污染报告"]
            values = {contamination.cell(row, 1).value: contamination.cell(row, 2).value for row in range(1, contamination.max_row + 1)}
            self.assertEqual(values["总Cluster数"], "5")
            self.assertEqual(values["状态亚群(非独立)"], "2")
            self.assertEqual(values["明确污染需去除"], "1")

            manifest = [json.loads(line) for line in (folder / "_script_manifest.jsonl").read_text().splitlines()]
            self.assertEqual(manifest[-1]["status"], "ok")
            self.assertEqual(manifest[-1]["expected_clusters"], 5)
            self.assertEqual(manifest[-1]["writer"], "styled")


if __name__ == "__main__":
    unittest.main()
