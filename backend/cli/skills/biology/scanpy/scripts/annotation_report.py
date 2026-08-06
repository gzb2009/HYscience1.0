#!/usr/bin/env python3
"""
Write a multi-sheet annotation workbook from a JSON annotation table.

Gold-standard layout (matches PDAC_TNK_annotation-style deliverables):
  1. Subset summary — every cluster, confidence ✓/⚠️/△/✗
  2. Keep for downstream — priority-ranked keep list
  3. Contamination report — exclude clusters + footer stats
  4. Key findings — titled numbered discoveries

Usage:
  python annotation_report.py --input annotations.json --output TNK_annotation.xlsx --locale zh
  python annotation_report.py --input annotations.json --output TNK_annotation.xlsx --locale en

Sheet / header language follows --locale. Gene symbols, cluster ids, and English
names stay as provided (do not translate).
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone


PACKS = {
    "en": {
        "summary": "{subset} subset summary",
        "keep": "Keep for downstream",
        "contamination": "Contamination report",
        "findings": "Key findings & recommendations",
        "summary_cols": [
            "Cluster",
            "Gene name",
            "Cell type",
            "Category",
            "Confidence",
            "Key top markers",
            "Notes",
            "Action",
        ],
        "keep_title": "Subsets recommended for downstream analysis",
        "keep_cols": ["Priority", "Cluster", "Gene name", "Cell type", "Analysis note"],
        "cont_title": "Contamination / abnormal cluster report",
        "cont_cols": [
            "Cluster",
            "Name",
            "Contamination type",
            "Foreign markers",
            "Foreign marker pct range",
            "Native lineage markers (if any)",
            "Action",
        ],
        "stats_label": "Stats",
        "stat_keys": [
            ("total", "Total clusters"),
            ("true_cells", "True lineage cells"),
            ("state", "State-only (non-independent)"),
            ("exclude", "Exclude (contamination)"),
            ("keep_downstream", "Keep for downstream"),
            ("contamination_rate", "Contamination rate"),
        ],
        "findings_title": "Key findings and follow-up recommendations",
        "no_findings": "No reportable finding was identified from the supplied annotation evidence.",
        "label_priority": "Label",
        "default_subset": "Cluster",
    },
    "zh": {
        "summary": "{subset}亚群汇总",
        "keep": "下游分析保留列表",
        "contamination": "污染报告",
        "findings": "关键发现与建议",
        "summary_cols": [
            "Cluster",
            "基因命名",
            "细胞类型",
            "分类",
            "可信度",
            "关键Top Markers",
            "备注",
            "建议",
        ],
        "keep_title": "建议保留用于下游分析的亚群",
        "keep_cols": ["优先级", "Cluster", "基因命名", "细胞类型", "分析建议"],
        "cont_title": "污染/异常Cluster详细报告",
        "cont_cols": [
            "Cluster",
            "命名",
            "污染类型",
            "关键外来Marker",
            "外来marker pct范围",
            "原谱系marker(若有)",
            "处理建议",
        ],
        "stats_label": "统计",
        "stat_keys": [
            ("total", "总Cluster数"),
            ("true_cells", "真正目标谱系细胞"),
            ("state", "状态亚群(非独立)"),
            ("exclude", "明确污染需去除"),
            ("keep_downstream", "建议保留用于下游分析"),
            ("contamination_rate", "污染率"),
        ],
        "findings_title": "亚群分析——关键发现与后续建议",
        "no_findings": "基于当前注释证据，未识别到可报告的关键发现。",
        "label_priority": "标注",
        "default_subset": "亚群",
    },
}


def locale_pack(locale: str):
    if locale.lower().startswith("zh"):
        return PACKS["zh"]
    return PACKS["en"]


class ValidationError(ValueError):
    def __init__(self, messages: list[str]):
        super().__init__("; ".join(messages))
        self.messages = messages


def text_of(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(text_of(item) for item in value if text_of(item))
    if isinstance(value, dict):
        return ", ".join(f"{key}: {text_of(item)}" for key, item in value.items())
    return str(value).strip()


def markers_of(value) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        parts = []
        for item in value:
            if not isinstance(item, dict):
                parts.append(text_of(item))
                continue
            gene = text_of(item.get("gene") or item.get("symbol") or item.get("name"))
            pct = text_of(item.get("pct") or item.get("percentage"))
            parts.append(f"{gene}({pct})" if gene and pct else gene or text_of(item))
        return ", ".join(part for part in parts if part)
    if isinstance(value, dict):
        return ", ".join(f"{gene}({text_of(pct)})" for gene, pct in value.items())
    return text_of(value)


def canonical_id(value) -> str:
    raw = text_of(value)
    if len(raw) > 1 and raw[:1].lower() == "c" and raw[1:].isdigit():
        return str(int(raw[1:]))
    if raw.isdigit():
        return str(int(raw))
    return raw.lower()


def normalize_cluster(cluster: dict) -> dict:
    item = dict(cluster)
    item["id"] = text_of(cluster.get("id") or cluster.get("cluster") or cluster.get("group"))
    item["name"] = text_of(cluster.get("name") or cluster.get("gene_name"))
    item["cell_type"] = text_of(cluster.get("cell_type") or cluster.get("cellType") or cluster.get("lineage"))
    item["category"] = text_of(cluster.get("category") or cluster.get("lineage"))
    item["markers"] = markers_of(cluster.get("markers") or cluster.get("features"))
    item["notes"] = text_of(cluster.get("notes") or cluster.get("note") or cluster.get("rationale"))
    item["action"] = text_of(cluster.get("action") or cluster.get("suggestion"))
    item["confidence"] = confidence_of(item)
    return item


def source_cluster_ids(path: str) -> list[str]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(8192)
        f.seek(0)
        delimiter = "\t" if path.lower().endswith(".tsv") else ","
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
        except csv.Error:
            pass
        reader = csv.DictReader(f, delimiter=delimiter)
        fields = reader.fieldnames or []
        aliases = {"cluster", "cluster_id", "clusterid", "group", "leiden", "seurat_clusters", "id"}
        column = next((field for field in fields if field.strip().lower() in aliases), None)
        if not column:
            raise ValidationError([f"markers source has no cluster column; found: {', '.join(fields) or '(none)'}"])
        values = [text_of(row.get(column)) for row in reader]
    ids = list(dict.fromkeys(value for value in values if value))
    if not ids:
        raise ValidationError(["markers source contains no cluster ids"])
    return ids


def normalize_data(data: dict, expected: list[str] | None = None) -> tuple[dict, dict]:
    raw = data.get("clusters")
    if not isinstance(raw, list) or not raw:
        raise ValidationError(["clusters must be a non-empty array"])

    clusters = [normalize_cluster(cluster) for cluster in raw if isinstance(cluster, dict)]
    errors = []
    if len(clusters) != len(raw):
        errors.append("every clusters item must be an object")

    required = ("id", "name", "cell_type", "category", "confidence", "markers", "notes", "action")
    for index, cluster in enumerate(clusters):
        missing = [key for key in required if not text_of(cluster.get(key))]
        if missing:
            errors.append(f"cluster[{index}] {cluster.get('id') or '(no id)'} missing: {', '.join(missing)}")
        if cluster.get("confidence") not in ("✓", "⚠️", "△", "✗"):
            errors.append(
                f"cluster[{index}] {cluster.get('id') or '(no id)'} confidence must be one of ✓/⚠️/△/✗"
            )

    ids = [canonical_id(cluster["id"]) for cluster in clusters]
    duplicates = sorted({item for item in ids if ids.count(item) > 1})
    if duplicates:
        errors.append(f"duplicate cluster ids: {', '.join(duplicates)}")

    expected_ids = [canonical_id(item) for item in expected or []]
    missing_ids = sorted(set(expected_ids) - set(ids))
    extra_ids = sorted(set(ids) - set(expected_ids)) if expected is not None else []
    if missing_ids:
        errors.append(f"missing cluster ids from markers source: {', '.join(missing_ids)}")
    if extra_ids:
        errors.append(f"cluster ids absent from markers source: {', '.join(extra_ids)}")

    contamination = data.get("contamination") or []
    if not isinstance(contamination, list):
        errors.append("contamination must be an array")
        contamination = []
    known = set(ids)
    for item in contamination:
        if not isinstance(item, dict):
            errors.append("every contamination item must be an object")
            continue
        cid = canonical_id(item.get("id"))
        if cid not in known:
            errors.append(f"contamination cluster absent from clusters: {text_of(item.get('id')) or '(no id)'}")
            continue
        cluster = clusters[ids.index(cid)]
        if cluster["confidence"] != "✗":
            errors.append(f"contamination cluster must have confidence ✗: {cluster['id']}")

    audit = {
        "expected_clusters": len(expected_ids) if expected is not None else None,
        "written_clusters": len(clusters),
        "missing_ids": missing_ids,
        "extra_ids": extra_ids,
        "duplicate_ids": duplicates,
    }
    if errors:
        raise ValidationError(errors)
    result = dict(data)
    result["clusters"] = clusters
    result["contamination"] = contamination
    return result, audit


def write_xlsx(path: str, sheets: dict[str, list[list[str]]]) -> str:
    """Write a styled workbook when openpyxl is available; fall back to a plain stdlib XLSX."""
    try:
        write_xlsx_styled(path, sheets)
        return "styled"
    except ModuleNotFoundError:
        write_xlsx_plain(path, sheets)
        return "plain"


def write_xlsx_styled(path: str, sheets: dict[str, list[list[str]]]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    title_fill = PatternFill("solid", fgColor="1F2937")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    sub_fill = PatternFill("solid", fgColor="E8F1FA")
    stat_fill = PatternFill("solid", fgColor="F3F4F6")
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, color="FFFFFF", size=13)
    bold = Font(bold=True)

    confidence_fill = {
        "✓": PatternFill("solid", fgColor="D9EAD3"),
        "⚠️": PatternFill("solid", fgColor="FFF2CC"),
        "⚠": PatternFill("solid", fgColor="FFF2CC"),
        "△": PatternFill("solid", fgColor="D9EAF7"),
        "✗": PatternFill("solid", fgColor="F4CCCC"),
    }
    priority_fill = {
        "★★★": PatternFill("solid", fgColor="FCE4D6"),
        "★★": PatternFill("solid", fgColor="FFF2CC"),
        "★": PatternFill("solid", fgColor="E2F0D9"),
        "⚠️": PatternFill("solid", fgColor="FFF2CC"),
        "⚠": PatternFill("solid", fgColor="FFF2CC"),
        "标注": PatternFill("solid", fgColor="DDEBF7"),
        "△": PatternFill("solid", fgColor="DDEBF7"),
    }

    def widths(cols: int) -> list[int]:
        if cols == 8:
            return [10, 28, 24, 16, 10, 58, 72, 46]
        if cols == 7:
            return [10, 28, 24, 56, 18, 46, 58]
        if cols == 5:
            return [10, 12, 28, 24, 64]
        if cols == 2:
            return [34, 105]
        return [18] * cols

    def title_row(rows: list[list[str]]) -> bool:
        return bool(rows and rows[0] and rows[0][0] and all(not x for x in rows[0][1:]))

    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        for row in rows:
            ws.append(row)

        max_col = max((len(r) for r in rows), default=1)
        max_row = len(rows)
        for i, width in enumerate(widths(max_col), start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A4" if title_row(rows) and max_row >= 3 else "A2"

        header = 3 if title_row(rows) and max_row >= 3 else 1
        if max_row >= header and max_col > 1:
            ws.auto_filter.ref = f"A{header}:{get_column_letter(max_col)}{max_row}"

        if title_row(rows):
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
            cell = ws.cell(1, 1)
            cell.fill = title_fill
            cell.font = title_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 24

        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = wrap
                if cell.row == header:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center

        # Findings sheet is a reading document, not a data table.
        if max_col == 2 and title_row(rows):
            ws.auto_filter.ref = None
            ws.freeze_panes = "A3"
            for r in range(3, max_row + 1):
                first = str(ws.cell(r, 1).value or "")
                if not first:
                    continue
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
                cell = ws.cell(r, 1)
                if first[:1].isdigit() and "." in first[:4]:
                    cell.fill = sub_fill
                    cell.font = bold
                    ws.row_dimensions[r].height = 22
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    ws.row_dimensions[r].height = max(34, min(240, 18 + len(first) // 48 * 15))
            continue

        for r in range(1, max_row + 1):
            if r == header:
                ws.row_dimensions[r].height = 24
            elif r != 1 or not title_row(rows):
                lengths = [
                    max((len(line) for line in str(ws.cell(r, c).value or "").splitlines()), default=0)
                    for c in range(1, max_col + 1)
                ]
                lines = max(1, max((length + 31) // 32 for length in lengths))
                ws.row_dimensions[r].height = min(180, max(24, lines * 18))
            values = [ws.cell(r, c).value for c in range(1, max_col + 1)]
            first = str(values[0] or "")
            if first in ("统计", "Stats"):
                for c in range(1, max_col + 1):
                    ws.cell(r, c).fill = stat_fill
                    ws.cell(r, c).font = bold
                ws.row_dimensions[r].height = 22

            for c in range(1, max_col + 1):
                value = str(ws.cell(r, c).value or "").strip()
                if value in confidence_fill:
                    ws.cell(r, c).fill = confidence_fill[value]
                    ws.cell(r, c).font = bold
                    ws.cell(r, c).alignment = center
                if value in priority_fill:
                    ws.cell(r, c).fill = priority_fill[value]
                    ws.cell(r, c).font = bold
                    ws.cell(r, c).alignment = center

        # Keep identifiers compact and long evidence readable.
        for c in (1, 2, 4, 5):
            if c <= max_col:
                for r in range(1, max_row + 1):
                    ws.cell(r, c).alignment = center if r == header or c in (1, 5) else wrap

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)


def write_xlsx_plain(path: str, sheets: dict[str, list[list[str]]]):
    """Minimal XLSX writer (stdlib fallback)."""
    import zipfile
    from xml.sax.saxutils import escape

    def col_name(i: int) -> str:
        s = ""
        n = i
        while True:
            s = chr(ord("A") + (n % 26)) + s
            n = n // 26 - 1
            if n < 0:
                break
        return s

    def sheet_xml(rows: list[list[str]]) -> str:
        lines = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
            "<sheetData>",
        ]
        for r, row in enumerate(rows, start=1):
            cells = []
            for c, val in enumerate(row):
                ref = f"{col_name(c)}{r}"
                text = escape("" if val is None else str(val))
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>')
            lines.append(f'<row r="{r}">{"".join(cells)}</row>')
        lines += ["</sheetData>", "</worksheet>"]
        return "\n".join(lines)

    names = list(sheets.keys())
    workbook = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
        "<sheets>",
    ]
    for i, name in enumerate(names, start=1):
        workbook.append(f'<sheet name="{escape(name)}" sheetId="{i}" r:id="rId{i}"/>')
    workbook += ["</sheets>", "</workbook>"]

    rels = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
    ]
    for i in range(1, len(names) + 1):
        rels.append(
            f'<Relationship Id="rId{i}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            f'Target="worksheets/sheet{i}.xml"/>'
        )
    rels.append("</Relationships>")

    content_types = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        '<Default Extension="xml" ContentType="application/xml"/>',
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>',
    ]
    for i in range(1, len(names) + 1):
        content_types.append(
            f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        )
    content_types.append("</Types>")

    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>
"""

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "\n".join(content_types))
        z.writestr("_rels/.rels", root_rels)
        z.writestr("xl/workbook.xml", "\n".join(workbook))
        z.writestr("xl/_rels/workbook.xml.rels", "\n".join(rels))
        for i, name in enumerate(names, start=1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", sheet_xml(sheets[name]))


def confidence_of(c: dict) -> str:
    raw = str(c.get("confidence") or c.get("flag") or "").strip()
    aliases = {
        "✓": "✓",
        "keep": "✓",
        "pass": "✓",
        "⚠": "⚠️",
        "⚠️": "⚠️",
        "warning": "⚠️",
        "caution": "⚠️",
        "keep-with-caution": "⚠️",
        "△": "△",
        "state": "△",
        "state-only": "△",
        "✗": "✗",
        "exclude": "✗",
        "contamination": "✗",
    }
    if raw.lower() in aliases:
        return aliases[raw.lower()]
    if c.get("exclude") or str(c.get("category", "")).lower() in ("contamination", "污染"):
        return "✗"
    if "cycl" in str(c.get("category", "")).lower() or "增殖" in str(c.get("category", "")):
        return "△"
    if "ifn" in str(c.get("category", "")).lower() or "IFN" in str(c.get("category", "")):
        return "△"
    if raw:
        return raw
    if c.get("keep") is True:
        return "✓"
    return ""


def priority_rank(p: str) -> int:
    order = {"★★★": 0, "★★": 1, "★": 2, "⚠️": 3, "⚠": 3, "标注": 4, "Label": 4, "△": 4}
    return order.get(str(p).strip(), 9)


def build_sheets(data: dict, pack: dict) -> dict[str, list[list[str]]]:
    subset = str(data.get("subset") or pack["default_subset"]).strip() or pack["default_subset"]
    clusters = data.get("clusters") or []
    contamination = data.get("contamination") or []
    findings = data.get("findings") or []
    stats = data.get("stats") or {}

    summary_name = pack["summary"].format(subset=subset)
    summary_rows: list[list[str]] = [pack["summary_cols"]]
    for c in clusters:
        conf = confidence_of(c)
        summary_rows.append(
            [
                text_of(c.get("id")),
                text_of(c.get("name")),
                text_of(c.get("cell_type") or c.get("cellType") or c.get("lineage")),
                text_of(c.get("category") or c.get("lineage")),
                conf,
                markers_of(c.get("markers") or c.get("features")),
                text_of(c.get("notes") or c.get("note")),
                text_of(c.get("action") or c.get("suggestion")),
            ]
        )

    keep_rows: list[list[str]] = [
        [pack["keep_title"], "", "", "", ""],
        ["", "", "", "", ""],
        pack["keep_cols"],
    ]
    keepers = []
    for c in clusters:
        conf = confidence_of(c)
        if conf == "✗":
            continue
        if c.get("keep") is False:
            continue
        keepers.append(c)
    keepers.sort(key=lambda c: (priority_rank(str(c.get("priority") or "")), str(c.get("id", ""))))
    for c in keepers:
        keep_rows.append(
            [
                text_of(c.get("priority") or (pack["label_priority"] if confidence_of(c) == "△" else "★")),
                text_of(c.get("id")),
                text_of(c.get("name")),
                text_of(c.get("cell_type") or c.get("cellType") or c.get("lineage")),
                text_of(c.get("analysis") or c.get("action")),
            ]
        )

    cont_rows: list[list[str]] = [
        [pack["cont_title"], "", "", "", "", "", ""],
        ["", "", "", "", "", "", ""],
        pack["cont_cols"],
    ]
    details = {canonical_id(c.get("id")): c for c in contamination if isinstance(c, dict)}
    cont_src = []
    for cluster in clusters:
        if confidence_of(cluster) != "✗":
            continue
        detail = details.get(canonical_id(cluster.get("id")), {})
        cont_src.append({**cluster, **detail, "id": cluster.get("id"), "name": cluster.get("name")})
    for c in cont_src:
        cont_rows.append(
            [
                text_of(c.get("id")),
                text_of(c.get("name")),
                text_of(c.get("kind") or c.get("type") or c.get("category") or "contamination"),
                markers_of(c.get("foreign_markers") or c.get("markers") or c.get("features")),
                text_of(c.get("pct_range") or c.get("pct")),
                markers_of(c.get("native_markers") or c.get("native")),
                text_of(c.get("action") or c.get("note")),
            ]
        )

    cont_rows.append(["", "", "", "", "", "", ""])
    cont_rows.append([pack["stats_label"], "", "", "", "", "", ""])
    n_ex = len(cont_src)
    n_true = sum(1 for c in clusters if confidence_of(c) in ("✓", "⚠️"))
    n_state = sum(1 for c in clusters if confidence_of(c) == "△")
    rate = f"{n_ex}/{len(clusters)} = {n_ex / len(clusters):.1%}" if clusters else "0/0"
    derived = {
        "total": len(clusters),
        "true_cells": n_true,
        "state": n_state,
        "exclude": n_ex,
        "keep_downstream": len(keepers),
        "contamination_rate": rate,
    }
    for key, value in derived.items():
        supplied = stats.get(key)
        if supplied not in (None, "") and text_of(supplied) != text_of(value):
            print(f"WARN: stats.{key}={supplied!r} differs from derived value {value!r}; using derived value", file=sys.stderr)
    for key, label in pack["stat_keys"]:
        cont_rows.append([label, str(derived.get(key, "")), "", "", "", "", ""])

    find_rows: list[list[str]] = [
        [str(data.get("findings_title") or pack["findings_title"]), ""],
        ["", ""],
    ]
    for item in findings:
        if isinstance(item, dict):
            title = str(item.get("title") or "")
            body = str(item.get("body") or item.get("text") or "")
            if title:
                find_rows.append([title, ""])
            if body:
                find_rows.append([body, ""])
        else:
            find_rows.append([str(item), ""])
    if len(find_rows) == 2:
        find_rows.append([pack["no_findings"], ""])

    return {
        summary_name: summary_rows,
        pack["keep"]: keep_rows,
        pack["contamination"]: cont_rows,
        pack["findings"]: find_rows,
    }


def write_manifest(path: str, manifest: dict):
    side = os.path.join(os.path.dirname(os.path.abspath(path)) or ".", "_script_manifest.jsonl")
    os.makedirs(os.path.dirname(side), exist_ok=True)
    with open(side, "a", encoding="utf-8") as f:
        f.write(json.dumps(manifest, ensure_ascii=False) + "\n")


def main():
    p = argparse.ArgumentParser(description="Build annotation xlsx from JSON")
    p.add_argument("--input", required=True, help="annotations.json")
    p.add_argument("--output", required=True, help="output .xlsx path")
    p.add_argument("--markers", help="source markers CSV/TSV used to verify the complete cluster set")
    p.add_argument("--locale", default="en", help="UI locale (en / zh / …)")
    args = p.parse_args()

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    out = args.output
    if not out.lower().endswith(".xlsx"):
        out = out + ".xlsx"
    base = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "pipeline": "scanpy-annotation-report",
        "output": os.path.abspath(out),
        "locale": args.locale,
    }
    try:
        expected = source_cluster_ids(args.markers) if args.markers else None
        data, audit = normalize_data(data, expected)
    except ValidationError as error:
        write_manifest(out, {**base, "status": "error", "errors": error.messages})
        for message in error.messages:
            print(f"ERROR: {message}", file=sys.stderr)
        return 2

    pack = locale_pack(args.locale)
    sheets = build_sheets(data, pack)
    writer = write_xlsx(out, sheets)
    clusters = data["clusters"]
    manifest = {
        **base,
        **audit,
        "status": "ok",
        "n_clusters": len(clusters),
        "markers_source": os.path.abspath(args.markers) if args.markers else None,
        "writer": writer,
        "sheets": list(sheets.keys()),
    }
    write_manifest(out, manifest)

    print(f"OK: wrote {out} ({len(clusters)} clusters; sheets={list(sheets.keys())})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
